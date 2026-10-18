from __future__ import annotations

from io import BytesIO
from pathlib import Path

import pytest

from app.services.document_extractor import (
    EmptyDocumentError,
    MAX_SOURCE_CHARS,
    SourceTextTooLongError,
    UnsupportedDocumentError,
    extract_document,
)


FIXTURES = Path(__file__).parent / "fixtures"


def test_markdown_decodes_bom_and_normalizes_newlines():
    result = extract_document("lesson.md", b"\xef\xbb\xbf# Lesson\r\n\r\nText")

    assert result.source_type == "md"
    assert result.text == "# Lesson\n\nText"


def test_markdown_preserves_headings_tables_lists_and_code_fences():
    source = """# Topic

| Name | Value |
| --- | --- |
| A | 1 |

- first
- second

```python
print('hello')
```
"""

    result = extract_document("lesson.md", source.encode())

    for fragment in ("# Topic", "| Name | Value |", "- first", "```python", "print('hello')"):
        assert fragment in result.text


def test_markdown_removes_html_comments_and_redundant_blank_lines():
    result = extract_document(
        "lesson.md", b"# Topic\n\n\n<!-- hidden note -->\n\n\nVisible text\n\n\n"
    )

    assert "hidden note" not in result.text
    assert "Visible text" in result.text
    assert "\n\n\n" not in result.text


def test_markdown_preserves_comment_like_text_inside_code_fence():
    result = extract_document("lesson.md", b"```html\n<!-- code example -->\n```\n")

    assert "<!-- code example -->" in result.text


def test_markdown_preserves_comment_like_text_inside_inline_and_indented_code():
    source = (
        "Use `<!-- inline example -->` in the template.\n\n"
        "    <!-- indented example -->\n"
    )

    result = extract_document("lesson.md", source.encode())

    assert "`<!-- inline example -->`" in result.text
    assert "    <!-- indented example -->" in result.text


def test_txt_decodes_utf8_bom():
    result = extract_document("notes.txt", b"\xef\xbb\xbfFirst line\r\nSecond line")

    assert result.source_type == "txt"
    assert result.text == "First line\nSecond line"


def test_docx_extracts_paragraphs_and_table_cells():
    from docx import Document

    document = Document()
    document.add_paragraph("Paragraph one")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Cell A"
    table.cell(0, 1).text = "Cell B"
    data = BytesIO()
    document.save(data)

    result = extract_document("notes.docx", data.getvalue())

    assert result.source_type == "docx"
    assert "Paragraph one" in result.text
    assert "Cell A" in result.text
    assert "Cell B" in result.text


def test_pdf_extracts_text_with_page_locator():
    result = extract_document("lesson.pdf", (FIXTURES / "tiny-text.pdf").read_bytes())

    assert result.source_type == "pdf"
    assert "PDF learning text for extraction." in result.text
    assert result.locators
    assert result.locators[0].label == "Page 1"


def test_rejects_unsupported_doc_extension():
    with pytest.raises(UnsupportedDocumentError):
        extract_document("legacy.doc", b"not a supported document")


def test_rejects_empty_or_scanned_pdf():
    from pypdf import PdfWriter

    writer = PdfWriter()
    writer.add_blank_page(width=612, height=792)
    data = BytesIO()
    writer.write(data)

    with pytest.raises(EmptyDocumentError):
        extract_document("scanned.pdf", data.getvalue())


def test_rejects_file_over_20mb_before_parsing():
    from app.services.document_extractor import MAX_UPLOAD_BYTES, UploadTooLargeError

    with pytest.raises(UploadTooLargeError):
        extract_document("lesson.md", b"x" * (MAX_UPLOAD_BYTES + 1))


def test_accepts_text_at_200000_char_limit():
    assert MAX_SOURCE_CHARS == 200_000
    result = extract_document("lesson.txt", ("x" * 200_000).encode())

    assert len(result.text) == 200_000


def test_rejects_text_over_200000_chars_without_truncation():
    with pytest.raises(SourceTextTooLongError):
        extract_document("lesson.txt", ("x" * 200_001).encode())


def _append_xml_text(parent, text: str) -> None:
    from docx.oxml import OxmlElement

    run = OxmlElement("w:r")
    text_element = OxmlElement("w:t")
    text_element.text = text
    run.append(text_element)
    parent.append(run)


def _append_wrapped_paragraph(parent, text: str, *, structured: bool = False) -> None:
    from docx.oxml import OxmlElement

    container = parent
    if structured:
        content = OxmlElement("w:sdtContent")
        parent.append(content)
        container = content
    paragraph = OxmlElement("w:p")
    _append_xml_text(paragraph, text)
    container.append(paragraph)


def test_docx_extracts_paragraphs_inside_word_block_wrappers():
    from docx import Document
    from docx.oxml import OxmlElement

    from app.services.extractors.docx import extract_docx

    document = Document()
    body = document.element.body
    for tag, text in (
        ("w:customXml", "custom xml text"),
        ("w:ins", "inserted text"),
        ("w:sdt", "content control text"),
        ("w:smartTag", "smart tag text"),
    ):
        wrapper = OxmlElement(tag)
        _append_wrapped_paragraph(wrapper, text, structured=tag == "w:sdt")
        body.insert(len(body) - 1, wrapper)

    # An incomplete content control should be ignored instead of breaking extraction.
    body.insert(len(body) - 1, OxmlElement("w:sdt"))
    data = BytesIO()
    document.save(data)

    text = extract_docx(data.getvalue())

    assert text.split("\n\n") == [
        "custom xml text",
        "inserted text",
        "content control text",
        "smart tag text",
    ]


def test_docx_extracts_inline_wrappers_hyperlinks_and_field_text():
    from docx import Document
    from docx.oxml import OxmlElement

    from app.services.extractors.docx import extract_docx

    document = Document()
    paragraph = document.add_paragraph("Before ")
    structured = OxmlElement("w:sdt")
    content = OxmlElement("w:sdtContent")
    structured.append(content)
    _append_xml_text(content, "inside ")
    paragraph._p.append(structured)

    inserted = OxmlElement("w:ins")
    _append_xml_text(inserted, "inserted ")
    paragraph._p.append(inserted)

    hyperlink = OxmlElement("w:hyperlink")
    _append_xml_text(hyperlink, "linked ")
    paragraph._p.append(hyperlink)

    field = OxmlElement("w:fldSimple")
    _append_xml_text(field, "field text")
    paragraph._p.append(field)
    data = BytesIO()
    document.save(data)

    assert extract_docx(data.getvalue()) == "Before inside inserted linked field text"


def test_docx_extracts_nested_table_cells_and_skips_empty_rows():
    from docx import Document

    from app.services.extractors.docx import extract_docx

    document = Document()
    table = document.add_table(rows=3, cols=2)
    table.cell(0, 0).text = "  Top left  "
    table.cell(0, 1).text = "Top right"
    nested = table.cell(1, 0).add_table(rows=1, cols=1)
    nested.cell(0, 0).text = "Nested cell"
    table.cell(1, 1).text = "Outer cell"
    data = BytesIO()
    document.save(data)

    text = extract_docx(data.getvalue())

    assert text.split("\n\n") == ["Top left | Top right", "Nested cell | Outer cell"]


def test_markdown_removes_multiline_comments_but_keeps_text_after_closing_marker():
    from app.services.extractors.markdown import extract_markdown

    text, locators = extract_markdown(
        b"Visible before <!-- hidden across lines\ncontinued hidden --> visible after"
    )

    assert text == "Visible before \n visible after"
    assert locators == [("Section 1", 0, len(text))]


def test_markdown_fences_close_only_with_matching_marker_and_length():
    from app.services.extractors.markdown import extract_markdown

    text, _locators = extract_markdown(
        b"~~~python\n<!-- kept in code -->\n```\n<!-- also kept -->\n~~~~\n<!-- removed -->\nVisible"
    )

    assert "<!-- kept in code -->" in text
    assert "<!-- also kept -->" in text
    assert "<!-- removed -->" not in text
    assert text.endswith("Visible")


def test_markdown_preserves_unmatched_inline_code_marker_while_removing_comments():
    from app.services.extractors.markdown import extract_markdown

    text, _locators = extract_markdown(b"Literal ` marker <!-- hidden --> and visible")

    assert text == "Literal ` marker  and visible"


def test_markdown_heading_locators_cover_sections_and_plain_text_fallback():
    from app.services.extractors.markdown import extract_markdown

    text, locators = extract_markdown(b"# First ##")
    plain_text, plain_locators = extract_markdown(b"No headings here")
    empty_text, empty_locators = extract_markdown(b"")

    assert [(label, text[start:end]) for label, start, end in locators] == [("Heading: First", "# First ##")]
    assert plain_text == "No headings here"
    assert plain_locators == [("Section 1", 0, len(plain_text))]
    assert empty_text == ""
    assert empty_locators == []
